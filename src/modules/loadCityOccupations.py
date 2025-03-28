import os
import json
import datetime
import pyautogui as ptg
from openpyxl import load_workbook
from unidecode import unidecode
import re
import time
import sys
from machine_utils import get_positions, get_image_path, get_windows_scale_factor, resize_image

def find_image_on_screen(image_path, confidence=0.8):
    """ Tenta localizar a imagem na tela com diferentes escalas """
    try:
        for scale in [1.0, 1.25, 1.5]:
            resized_img = resize_image(image_path, scale)
            if resized_img is None:
                continue
            location = ptg.locateOnScreen(resized_img, grayscale=True, confidence=confidence)
            if location:
                return location
    except ptg.ImageNotFoundException:
        return None
    return None

positions = get_positions("Occupation")
spreadsheetOccupationsPath = sys.argv[1]

# Capturar o nome do município
padrao = r"/ocupacoes_(.*?)(?=\.)"
resultado = re.search(padrao, spreadsheetOccupationsPath)
city = resultado.group(1)

imgOcupacaoJaCadatrada = get_image_path("ocupacaoJaCadastrada.png")

def click_and_wait(position, duration=0.5):
    ptg.moveTo(position, duration=duration)
    ptg.click()
    time.sleep(duration)

def click_and_write(position, text, duration=0.5):
    ptg.moveTo(position, duration=duration)
    ptg.click(position)
    ptg.write(text.strip())
    time.sleep(duration)    

try:
    scale_factor = get_windows_scale_factor()
    imgOcupacaoJaCadatrada = resize_image(imgOcupacaoJaCadatrada, scale_factor)
    
    spreadsheetOccupations = load_workbook(spreadsheetOccupationsPath)
    sheetOccupations = spreadsheetOccupations["Ocupacoes"]
    qtd_linhas = sheetOccupations.max_row

    totalOccupationsRead = 0
    totalOccupationsFound = 0
    totalOccupationsInserted = 0
    ptg.PAUSE = 0.3

    ptg.hotkey("alt", "tab", interval=0.2)
    time.sleep(1)
    click_and_wait(positions["selTipoOcupacao"])

    for linha in range(2, qtd_linhas + 1):
        totalOccupationsRead += 1

        click_and_wait(positions["btnCriarOcupacao"])
        click_and_wait(positions["btnSelecaoServidor"])

        click_and_write(positions["cmpNomeOcupacao"], unidecode(sheetOccupations[f"A{linha}"].value), 0.5)
        time.sleep(0.3)

        click_and_wait(positions["btnSalvarOcupacao"])
        ptg.moveTo(positions["posicaoForaTela"])

        if imgOcupacaoJaCadatrada:
            location = find_image_on_screen(imgOcupacaoJaCadatrada)
        else:
            location = None

        if location:
            totalOccupationsFound += 1
            click_and_wait(positions["btnFecharAlerta"])
            time.sleep(1)
            click_and_wait(positions["btnCancelarOcupacao"])
            continue

        totalOccupationsInserted += 1
        time.sleep(1)
    
    # Criar log JSON
    now = datetime.datetime.now().strftime("%d/%m/%Y - %H:%M:%S")
    log_data = {
        "Ocupaçoes": f"Carga de Ocupações do município {city} - em {now}",
        "Status": "Realizada com Sucesso!!!",
        "Total de Ocupações Lidas": totalOccupationsRead,
        "Total de Ocupações já Cadastradas": totalOccupationsFound,
        "Total de Ocupações Novas Cadastradas": totalOccupationsInserted
    }
    log_file_path = os.path.join(os.path.dirname(spreadsheetOccupationsPath), "occupation_log.json")
    with open(log_file_path, "w", encoding="utf-8") as log_file:
        json.dump(log_data, log_file, indent=4, ensure_ascii=False)

    print(f"SUCCESS|{totalOccupationsRead}|{totalOccupationsFound}|{totalOccupationsInserted}")

except Exception as e:
    print(f"ERROR|{str(e)}")
    sys.exit(1)

ptg.hotkey("alt", "tab", interval=0.2)
